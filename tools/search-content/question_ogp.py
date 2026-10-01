"""Search Question 用 OGP（1200 × 630）の生成（ACR-20261001-019）。
Brand OGP System（assets/ogp/master/OGP_Template_Master.pptx）の Slide 14（Dreamin' Spiral 🌱 Home）を土台に、
ブランドライン = DREAMIN' SPIRAL 🌱 ／ タイトル = Theme ／ サブコピー = Question として書き出す（背景 ・ ロゴ ・ 配色 ・ 余白 ・ フォントは変えない）。
必要：python-pptx ・ LibreOffice（soffice）・ pdftoppm。使い方：b8e-lab の root で python3 tools/search-content/question_ogp.py"""
import os, sys, subprocess, tempfile, shutil
import copy
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from questions_data import Q
MASTER = 'assets/ogp/master/OGP_Template_Master.pptx'
OUT = 'assets/ogp/questions'
BASE_SLIDE = 14
R_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'

def render(key, title, sub_lines):
    p = Presentation(MASTER)
    sl = p.slides[BASE_SLIDE - 1]
    tfs = [sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    for shape, lines in ((tfs[1], [title]), (tfs[2], sub_lines)):
        para = shape.text_frame.paragraphs[0]
        runs = para.runs
        for r in runs[1:]: r._r.getparent().remove(r._r)
        first = runs[0]._r
        first.find(qn('a:t')).text = lines[0]
        prev = first
        for line in lines[1:]:
            br = etree.SubElement(para._p, qn('a:br'))
            rpr = first.find(qn('a:rPr'))
            if rpr is not None: br.append(copy.deepcopy(rpr))
            prev.addnext(br)
            nr = copy.deepcopy(first); nr.find(qn('a:t')).text = line
            br.addnext(nr); prev = nr
    ids = p.slides._sldIdLst
    for i, sid in reversed(list(enumerate(list(ids)))):
        if i != BASE_SLIDE - 1:
            p.part.drop_rel(sid.get(R_NS)); ids.remove(sid)
    tmp = tempfile.mkdtemp()
    try:
        pp = os.path.join(tmp, key + '.pptx'); p.save(pp)
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, pp], check=True, capture_output=True)
        subprocess.run(['pdftoppm', '-png', '-singlefile', '-scale-to-x', '1200', '-scale-to-y', '630', os.path.join(tmp, key + '.pdf'), os.path.join(tmp, key)], check=True)
        os.makedirs(OUT, exist_ok=True)
        shutil.copy(os.path.join(tmp, key + '.png'), os.path.join(OUT, key + '.png'))
    finally:
        shutil.rmtree(tmp)

if __name__ == '__main__':
    for q in Q:
        render(q['slug'], q['theme'], q.get('ogp_lines', q['qh']))  # OGP は 2 行まで（3 行だと Title に近づく）
    render('questions-hub', '問いから読む', ['日々の中で、ふと気になったこと。'])
    print('rendered', len(Q) + 1)
